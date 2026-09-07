"""Inventario formal de fuentes del proyecto Saber 11 x IPM (solo mide y describe).

Proposito: perfilar los 11 CSV de datos/raw/ (uno por uno, en chunks, sin cargarlos
    completos en memoria) y el anexo IPM datos/anex-PMultidimensional-Departamental-2024.xlsx
    (solo lectura, sin transformar ni exportar nada del xlsx).
Entradas: datos/raw/Examen_Saber_11_*_limpio.csv (11 archivos, sep ';', utf-8) y
    datos/anex-PMultidimensional-Departamental-2024.xlsx.
Salidas: datos/output/reports/01_inventario.md, datos/output/reports/01_inventario_icfes.csv
    (una fila por periodo) y datos/output/reports/01_inventario_ipm.md.
Sesion de bitacora que lo origina: SESION-002 (2026-09-07, entorno e inventario).
Reglas: sin filtros ni imputaciones (solo diagnostico), un CSV en memoria a la vez,
    chunksize=250_000, rutas relativas con pathlib, determinista e idempotente
    (re-ejecutar sobrescribe los mismos reportes).

Formato: celdas estilo jupytext py:percent (markdown + codigo) para convertir a notebook.
"""

# %% [markdown]
# # 01 — Inventario de fuentes (SESION-002)
# Este script **solo mide y describe**: no aplica filtros, imputaciones ni tratamientos.
# Lee cada CSV del ICFES en `chunks` de 250 000 filas (`dtype=str`) y perfila el anexo
# IPM en modo solo lectura.

# %%
import hashlib
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# %% [markdown]
# ## Constantes y rutas relativas
# `COLUMNAS_ESPERADAS_59` es la lista literal verificada en SESION-001
# (`head -n 1 ... | tr ';' '\n' | nl` sobre los 11 archivos, cabeceras identicas).

# %%
RAIZ = Path(__file__).resolve().parents[1]
DATOS_RAW = RAIZ / "datos" / "raw"
XLSX_IPM = RAIZ / "datos" / "anex-PMultidimensional-Departamental-2024.xlsx"
DIR_SALIDA = RAIZ / "datos" / "output" / "reports"
CHUNKSIZE = 250_000
SEPARADOR = ";"
CODIFICACION = "utf-8"

COLUMNAS_ESPERADAS_59 = [
    "periodo", "cole_area_ubicacion", "cole_bilingue", "cole_cod_depto_ubicacion",
    "cole_cod_mcpio_ubicacion", "cole_depto_ubicacion", "cole_jornada",
    "cole_mcpio_ubicacion", "cole_naturaleza", "cole_sede_principal",
    "estu_agregado", "estu_cod_depto_presentacion", "estu_cod_mcpio_presentacion",
    "estu_cod_reside_depto", "estu_cod_reside_mcpio", "estu_dedicacioninternet",
    "estu_dedicacionlecturadiaria", "estu_depto_presentacion", "estu_depto_reside",
    "estu_discapacidad", "estu_fechanacimiento", "estu_genero", "estu_grado",
    "estu_horassemanatrabaja", "estu_inse_individual", "estu_mcpio_presentacion",
    "estu_mcpio_reside", "estu_nacionalidad", "estu_nse_establecimiento",
    "estu_nse_individual", "estu_pais_reside", "estu_repite", "estu_tiporemuneracion",
    "fami_comecarnepescadohuevo", "fami_comecerealfrutoslegumbre",
    "fami_comelechederivados", "fami_cuartoshogar", "fami_educacionmadre",
    "fami_educacionpadre", "fami_estratovivienda", "fami_numlibros",
    "fami_personashogar", "fami_situacioneconomica", "fami_tieneautomovil",
    "fami_tienecomputador", "fami_tieneconsolavideojuegos", "fami_tienehornomicroogas",
    "fami_tieneinternet", "fami_tienelavadora", "fami_tienemotocicleta",
    "fami_tieneserviciotv", "fami_trabajolabormadre", "fami_trabajolaborpadre",
    "punt_c_naturales", "punt_global", "punt_ingles", "punt_lectura_critica",
    "punt_matematicas", "punt_sociales_ciudadanas",
]
COLUMNA_DEPTO = "cole_cod_depto_ubicacion"
COLUMNA_MCPio = "cole_cod_mcpio_ubicacion"
COLUMNAS_PUNTAJE = {
    "punt_global": 500,
    "punt_matematicas": 100,
    "punt_lectura_critica": 100,
    "punt_c_naturales": 100,
    "punt_sociales_ciudadanas": 100,
    "punt_ingles": 100,
}
AREAS_5 = [
    "punt_matematicas", "punt_lectura_critica", "punt_c_naturales",
    "punt_sociales_ciudadanas", "punt_ingles",
]
# Cifras EXPECTADAS del documento del proyecto (regla 4 de AGENTS.md: comparacion,
# nunca resultados). No existen en el repo; llegan por la consigna de la sesion.
EXPECTATIVA_TOTAL_REGISTROS = 4_039_755
EXPECTATIVA_FILTROS_REF = (606_214, 22, 28_394)

# %% [markdown]
# ## Utilidades de perfilado ICFES
# Nulos: con `dtype=str`, los campos vacios llegan como `NaN` (deteccion NA por defecto
# de pandas) y se cuentan con `isna()`. Numericos: `pd.to_numeric(errors="coerce")`;
# los valores no numericos se cuentan como `na_coerce ∧ no_na_original` (diagnostico
# explicito, sin silencios). Duplicados: md5 por fila sobre los campos unidos con el
# separador de unidad `\x1f` (evita colisiones por `;` embebidos en texto); se guardan
# los primeros 8 bytes del digest como entero en un `set` por archivo.

# %%
def es_nulo(valor) -> bool:
    """Verdadero si el valor es NaN (campo vacio leido como str)."""
    return valor != valor


def fila_a_bytes(fila) -> bytes:
    """Serializa una fila a bytes para el hash md5 (NaN -> cadena vacia)."""
    return "\x1f".join("" if es_nulo(v) else str(v) for v in fila).encode("utf-8")


def hash_fila_8bytes(bloque: bytes) -> int:
    """Entero de 8 bytes con los primeros 8 bytes del digest md5."""
    return int.from_bytes(hashlib.md5(bloque).digest()[:8], "big")


def clasificar_formato_codigos(unicos: set) -> str:
    """Clasifica el formato de un conjunto de codigos leidos como texto."""
    if not unicos:
        return "vacio"
    son_digitos = [c.isdigit() for c in unicos]
    if not all(son_digitos):
        return "mixto_no_solo_digitos"
    con_cero = [c for c in unicos if len(c) > 1 and c.startswith("0")]
    if con_cero:
        return "texto_con_ceros_izquierda"
    return "digitos_sin_ceros_izquierda"


def fmt_num(valor) -> str:
    """Formatea un estadistico float de forma determinista."""
    if valor != valor:
        return "NA"
    return f"{valor:.2f}"


def periodo_esperado(nombre_archivo: str) -> str:
    """Extrae el periodo YYYYP del nombre Examen_Saber_11_YYYYP_limpio.csv."""
    coincidencia = re.search(r"(\d{5})", nombre_archivo)
    return coincidencia.group(1) if coincidencia else "DESCONOCIDO"


# %% [markdown]
# ## Perfilado de un CSV del ICFES (una pasada en chunks + segunda pasada opcional)
# Por cada archivo se acumulan: registros, unicos de `periodo`, estadisticos de los 6
# puntajes, nulos por columna, unicos de depto/mcpio, hashes md5 y conteos diagnostico
# de los filtros de Fase 1 (sin aplicarlos).

# %%
def perfilar_csv_icfes(ruta: Path) -> dict:
    """Perfila un CSV del ICFES en chunks. Devuelve el dict de resultados del archivo."""
    lector = pd.read_csv(
        ruta, sep=SEPARADOR, encoding=CODIFICACION, dtype=str, chunksize=CHUNKSIZE
    )
    registros = 0
    periodos_vistos: set = set()
    columnas_reales: list | None = None
    nulos_por_columna = {c: 0 for c in COLUMNAS_ESPERADAS_59}
    columnas_extra: dict = {}
    depto_unicos: set = set()
    mcpio_unicos: set = set()
    depto_nulos = 0
    punt_acum = {
        c: {"nulos": 0, "no_numericos": 0, "suma": 0.0, "n_validos": 0,
            "min": None, "max": None, "fuera_escala": 0}
        for c in COLUMNAS_PUNTAJE
    }
    diag_depto_nulo = 0
    diag_global_cero = 0
    diag_global_nulo = 0
    diag_alguna_area_nula = 0
    hashes_vistos: set = set()

    for chunk in lector:
        if columnas_reales is None:
            columnas_reales = list(chunk.columns)
        registros += len(chunk)
        periodos_vistos.update(chunk["periodo"].dropna().unique().tolist())
        for columna in columnas_reales:
            if columna in nulos_por_columna:
                nulos_por_columna[columna] += int(chunk[columna].isna().sum())
            else:
                columnas_extra[columna] = columnas_extra.get(columna, 0) + int(
                    chunk[columna].isna().sum()
                )
        serie_depto = chunk[COLUMNA_DEPTO]
        depto_nulos += int(serie_depto.isna().sum())
        depto_unicos.update(serie_depto.dropna().unique().tolist())
        mcpio_unicos.update(chunk[COLUMNA_MCPio].dropna().unique().tolist())
        for columna, limite in COLUMNAS_PUNTAJE.items():
            serie = chunk[columna]
            es_na = serie.isna()
            punt_acum[columna]["nulos"] += int(es_na.sum())
            numerica = pd.to_numeric(serie, errors="coerce")
            na_numerica = numerica.isna()
            punt_acum[columna]["no_numericos"] += int((na_numerica & ~es_na).sum())
            validos = numerica.dropna()
            punt_acum[columna]["n_validos"] += int(validos.shape[0])
            punt_acum[columna]["suma"] += float(validos.sum())
            if validos.shape[0] > 0:
                minimo = float(validos.min())
                maximo = float(validos.max())
                previo_min = punt_acum[columna]["min"]
                previo_max = punt_acum[columna]["max"]
                punt_acum[columna]["min"] = minimo if previo_min is None else min(previo_min, minimo)
                punt_acum[columna]["max"] = maximo if previo_max is None else max(previo_max, maximo)
                punt_acum[columna]["fuera_escala"] += int(((validos < 0) | (validos > limite)).sum())
        diag_depto_nulo += int(chunk[COLUMNA_DEPTO].isna().sum())
        serie_global = chunk["punt_global"]
        es_na_global = serie_global.isna()
        diag_global_nulo += int(es_na_global.sum())
        numerica_global = pd.to_numeric(serie_global, errors="coerce")
        diag_global_cero += int((numerica_global == 0).sum())
        diag_alguna_area_nula += int(chunk[AREAS_5].isna().any(axis=1).sum())
        for fila in chunk.itertuples(index=False, name=None):
            hashes_vistos.add(hash_fila_8bytes(fila_a_bytes(fila)))

    duplicados = registros - len(hashes_vistos)
    for columna, acum in punt_acum.items():
        acum["media"] = acum["suma"] / acum["n_validos"] if acum["n_validos"] > 0 else float("nan")
    return {
        "archivo": ruta.name,
        "periodo_esperado": periodo_esperado(ruta.name),
        "periodos_vistos": sorted(periodos_vistos),
        "registros": registros,
        "columnas_reales": columnas_reales if columnas_reales is not None else [],
        "columnas_identicas": (columnas_reales == COLUMNAS_ESPERADAS_59),
        "columnas_extra": columnas_extra,
        "nulos_por_columna": nulos_por_columna,
        "depto_nulos": depto_nulos,
        "depto_unicos": sorted(depto_unicos),
        "depto_formato": clasificar_formato_codigos(depto_unicos),
        "mcpio_nunique": len(mcpio_unicos),
        "puntajes": punt_acum,
        "duplicados": duplicados,
        "diag_depto_nulo": diag_depto_nulo,
        "diag_global_cero": diag_global_cero,
        "diag_global_nulo": diag_global_nulo,
        "diag_alguna_area_nula": diag_alguna_area_nula,
    }


def extraer_ejemplos_duplicados(ruta: Path, max_ejemplos: int = 3) -> list:
    """Segunda pasada sobre el CSV: captura hasta max_ejemplos filas duplicadas."""
    lector = pd.read_csv(
        ruta, sep=SEPARADOR, encoding=CODIFICACION, dtype=str, chunksize=CHUNKSIZE
    )
    vistos: set = set()
    ejemplos: list = []
    for chunk in lector:
        for fila in chunk.itertuples(index=False, name=None):
            marca = hash_fila_8bytes(fila_a_bytes(fila))
            if marca in vistos and len(ejemplos) < max_ejemplos:
                ejemplos.append(SEPARADOR.join("" if es_nulo(v) else str(v) for v in fila))
            else:
                vistos.add(marca)
        if len(ejemplos) >= max_ejemplos:
            break
    return ejemplos


# %% [markdown]
# ## Perfilado del anexo IPM (solo lectura)
# Se listan las hojas con dimensiones (openpyxl, sin guardar), se vuelcan las primeras
# 60 filas de `IPM_Departamentos` tal cual (`header=None`) y se analiza la hoja completa
# (pequena) para responder con cifras: anios, unidades por anio, filas no
# departamentales, columna DIVIPOLA y lista de columnas.

# %%
PALABRAS_CLAVE_NO_DEPTAL = {
    "total nacional", "total", "nacional", "cabeceras", "cabecera",
    "centros poblados y rural disperso", "rural", "urbano",
    "caribe", "andina", "andino", "pacifica", "pacifico", "orinoquia",
    "orinoquia-amazonia", "amazonia", "central", "oriental", "occidental", "region",
}


def normalizar_texto(texto: str) -> str:
    """Minusculas sin tildes ni espacios sobrantes para comparar nombres."""
    import unicodedata

    base = unicodedata.normalize("NFD", str(texto).strip().lower())
    base = "".join(c for c in base if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", base)


def clasificar_territorio(nombre: str) -> str:
    """Clasifica un nombre territorial por coincidencia EXACTA normalizada.

    Coincidencia exacta (no subcadena) para no marcar 'Amazonas' como region
    'amazonia' ni 'Atlántico' como nada ajeno.
    """
    texto = normalizar_texto(nombre)
    if texto in PALABRAS_CLAVE_NO_DEPTAL:
        return f"no_departamental:{texto}"
    return "posible_departamento"


def extraer_anio(valor) -> str | None:
    """Devuelve el anio (cadena 'YYYY') si la celda lo contiene, o None.

    Cubre enteros/flotantes enteros en [2010, 2030] y textos con notas al pie
    tipo '2020**' o '2021***' (anexo DANE: notas de inasistencia y operativo ECV).
    """
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)) and valor == valor:
        if 2010 <= int(valor) <= 2030 and float(valor) == float(int(valor)):
            return str(int(valor))
        return None
    if isinstance(valor, str):
        coincidencia = re.search(r"(19|20)\d{2}", valor)
        if coincidencia and 2010 <= int(coincidencia.group(0)) <= 2030:
            return coincidencia.group(0)
    return None


def celdas_parecen_anio(valor) -> bool:
    """Verdadero si la celda contiene un anio plausible (2010-2030)."""
    return extraer_anio(valor) is not None


def perfilar_ipm() -> dict:
    """Perfila el xlsx del IPM. Devuelve hojas, volcado de 60 filas y analisis."""
    libro = load_workbook(XLSX_IPM, data_only=True)
    hojas = []
    for nombre in libro.sheetnames:
        hoja = libro[nombre]
        hojas.append({
            "hoja": nombre,
            "dimension": hoja.dimensions,
            "filas_max": hoja.max_row,
            "columnas_max": hoja.max_column,
        })
    df_muestra = pd.read_excel(
        XLSX_IPM, sheet_name="IPM_Departamentos", header=None, nrows=60, engine="openpyxl"
    )
    lineas_volcado = []
    for _, fila in df_muestra.iterrows():
        lineas_volcado.append("|".join("" if es_nulo(v) else str(v) for v in fila.tolist()))
    df_full = pd.read_excel(
        XLSX_IPM, sheet_name="IPM_Departamentos", header=None, engine="openpyxl"
    )
    # Fila de grupos de anio: la que contiene 'departamento' junto a celdas de anio
    # (layout ancho DANE: 1 columna territorio + 7 anios x 3 dominios). La fila
    # siguiente trae los sub-encabezados Total/Cabeceras/Rural por anio.
    fila_grupos = None
    for i, fila in df_full.iterrows():
        textos = [str(v).strip().lower() for v in fila.tolist() if not es_nulo(v)]
        anios_fila = {a for a in (extraer_anio(v) for v in fila.tolist()) if a is not None}
        # Se exigen >=2 anios distintos para no confundir el titulo combinado
        # ('Incidencia ... Departamentos 2018-2024' en una sola celda) con el encabezado.
        if any("departamento" in t for t in textos) and len(anios_fila) >= 2:
            fila_grupos = int(i)
            break
    fila_sub = fila_grupos + 1 if fila_grupos is not None else None
    posiciones_anio: list = []
    anios_en_orden: list = []
    columnas_detectadas: list = []
    df_datos = df_full
    if fila_grupos is not None:
        fila_anios = df_full.iloc[fila_grupos].tolist()
        ultimo_anio = None
        for valor in fila_anios:
            anio = extraer_anio(valor)
            if anio is not None:
                ultimo_anio = anio
            if ultimo_anio is not None:
                anios_en_orden.append(ultimo_anio)
        anios_detectados = sorted(set(anios_en_orden))
        posiciones_anio = [
            (fila_grupos, j, a) for j, a in enumerate(anios_en_orden) if a is not None
        ]
        fila_encabezado = fila_grupos
        sub_encabezados = [
            "" if es_nulo(v) else " ".join(str(v).split()) for v in df_full.iloc[fila_sub].tolist()
        ]
        primer_territorio = " ".join(str(fila_anios[0]).split()) if not es_nulo(fila_anios[0]) else "territorio"
        columnas_detectadas = [primer_territorio]
        # anios_en_orden ya excluye la columna 0 (territorio): va 1:1 con sub[1:].
        for anio, sub in zip(anios_en_orden, sub_encabezados[1:]):
            columnas_detectadas.append(f"{anio} {sub}".strip())
        # Filas de datos: desde fila_sub+1 hasta la primera fila de notas
        # (celda inicial vacia o que inicia con Fuente/Nota/Datos/Actualizado).
        inicio_datos = fila_sub + 1
        fin_datos = int(df_full.shape[0])
        for i in range(inicio_datos, int(df_full.shape[0])):
            primera = df_full.iloc[i, 0]
            if es_nulo(primera) or str(primera).strip() == "":
                fin_datos = i
                break
            if normalizar_texto(primera).split(" ")[0] in {"fuente", "nota", "notas", "datos", "actualizado"}:
                fin_datos = i
                break
        df_datos = df_full.iloc[inicio_datos:fin_datos].reset_index(drop=True)
        filas_con_anio = sorted({fila_grupos})
        columnas_con_anio = sorted({p[1] for p in posiciones_anio})
    else:
        anios_detectados = sorted({
            a for _, fila in df_full.iterrows() for a in
            [extraer_anio(v) for v in fila.tolist()] if a is not None
        })
        filas_con_anio = []
        columnas_con_anio = []
        fila_encabezado = None
    col_territorio = None
    col_codigo = None
    for j, nombre_col in enumerate(columnas_detectadas):
        bajo = nombre_col.lower()
        if col_territorio is None and any(
            k in bajo for k in ["departamento", "dominio", "territor", "dpto", "nombre"]
        ):
            col_territorio = j
        if col_codigo is None and any(
            k in bajo for k in ["cod", "dane", "divipola"]
        ):
            col_codigo = j
    if col_territorio is None:
        mejor_j, mejor_n = None, -1
        for j in range(df_datos.shape[1]):
            n_texto = sum(
                isinstance(v, str) and len(v.strip()) >= 3 for v in df_datos.iloc[:, j].tolist()
            )
            if n_texto > mejor_n:
                mejor_n, mejor_j = n_texto, j
        col_territorio = mejor_j
    territorios = [v for v in df_datos.iloc[:, col_territorio].tolist() if not es_nulo(v)]
    territorios_unicos = sorted({str(v).strip() for v in territorios})
    clasificacion = {}
    for nombre in territorios_unicos:
        clase = clasificar_territorio(nombre)
        clasificacion.setdefault(clase, []).append(nombre)
    codigos_divipola = None
    if col_codigo is not None:
        codigos_divipola = sorted({
            str(v).strip() for v in df_datos.iloc[:, col_codigo].tolist() if not es_nulo(v)
        })
    cobertura_por_columna = {
        int(j): int(df_datos.iloc[:, j].notna().sum()) for j in range(df_datos.shape[1])
    }
    return {
        "hojas": hojas,
        "n_hojas": len(hojas),
        "muestra_nfilas": int(df_muestra.shape[0]),
        "muestra_ncolumnas": int(df_muestra.shape[1]),
        "volcado_60": lineas_volcado,
        "full_filas": int(df_full.shape[0]),
        "full_columnas": int(df_full.shape[1]),
        "n_filas_datos": int(df_datos.shape[0]),
        "anios_detectados": anios_detectados,
        "filas_con_anio": filas_con_anio,
        "columnas_con_anio": columnas_con_anio,
        "fila_encabezado": fila_encabezado,
        "columnas_detectadas": columnas_detectadas,
        "cobertura_por_columna": cobertura_por_columna,
        "col_territorio": col_territorio,
        "col_codigo": col_codigo,
        "territorios_unicos": territorios_unicos,
        "n_territorios_unicos": len(territorios_unicos),
        "clasificacion": {k: sorted(v) for k, v in clasificacion.items()},
        "codigos_divipola": codigos_divipola,
    }


# %% [markdown]
# ## Escritura de artefactos y funcion principal
# Se generan los tres artefactos de la TAREA C. Los checkpoints de la TAREA D usan las
# expectativas del documento solo como referencia cruzada (regla 4).

# %%
def fila_resumen_csv(resultado: dict) -> dict:
    """Aplana el resultado de un periodo a una fila del CSV de inventario."""
    fila = {
        "archivo": resultado["archivo"],
        "periodo_esperado": resultado["periodo_esperado"],
        "periodos_vistos": "|".join(resultado["periodos_vistos"]),
        "registros": resultado["registros"],
        "n_columnas": len(resultado["columnas_reales"]),
        "columnas_identicas_59": resultado["columnas_identicas"],
        "columnas_extra": "|".join(sorted(resultado["columnas_extra"].keys())),
    }
    for columna, limite in COLUMNAS_PUNTAJE.items():
        acum = resultado["puntajes"][columna]
        fila[f"{columna}_nulos"] = acum["nulos"]
        fila[f"{columna}_no_numericos"] = acum["no_numericos"]
        fila[f"{columna}_min"] = acum["min"]
        fila[f"{columna}_max"] = acum["max"]
        fila[f"{columna}_media"] = acum["media"]
        fila[f"{columna}_fuera_escala_0_{limite}"] = acum["fuera_escala"]
    fila["depto_nulos"] = resultado["depto_nulos"]
    fila["depto_nunique"] = len(resultado["depto_unicos"])
    fila["depto_formato"] = resultado["depto_formato"]
    fila["depto_codigos"] = "|".join(resultado["depto_unicos"])
    fila["mcpio_nunique"] = resultado["mcpio_nunique"]
    fila["duplicados_exactos"] = resultado["duplicados"]
    fila["diag_depto_nulo"] = resultado["diag_depto_nulo"]
    fila["diag_global_cero"] = resultado["diag_global_cero"]
    fila["diag_global_nulo"] = resultado["diag_global_nulo"]
    fila["diag_alguna_area_nula"] = resultado["diag_alguna_area_nula"]
    fila["cole_bilingue_nulos"] = resultado["nulos_por_columna"].get("cole_bilingue", 0)
    nulos = {c: n for c, n in resultado["nulos_por_columna"].items() if n > 0}
    fila["nulos_por_columna_solo_mayor_0"] = "|".join(
        f"{c}={n}" for c, n in sorted(nulos.items(), key=lambda x: -x[1])
    )
    return fila


def escribir_inventario_md(resultados: list, total_registros: int) -> None:
    """Escribe datos/output/reports/01_inventario.md con tablas por archivo y checkpoints."""
    lineas = []
    lineas.append("# 01 — Inventario de fuentes ICFES + IPM (SESION-002)")
    lineas.append("")
    lineas.append("Solo mide y describe: sin filtros, imputaciones ni tratamientos.")
    lineas.append("Fuente de cifras: ejecucion de scripts/01_inventario.py (chunks de 250_000).")
    lineas.append("")
    lineas.append("## B1 — ICFES por archivo")
    lineas.append("")
    lineas.append("| archivo | periodo | registros | columnas | cabecera_59 | duplicados |")
    lineas.append("|---|---|---|---|---|---|")
    for r in resultados:
        lineas.append(
            f"| {r['archivo']} | {r['periodo_esperado']} | {r['registros']} | "
            f"{len(r['columnas_reales'])} | {r['columnas_identicas']} | {r['duplicados']} |"
        )
    lineas.append("")
    lineas.append(f"Total registros 11 archivos: **{total_registros}**.")
    lineas.append("")
    for r in resultados:
        lineas.append(f"### {r['periodo_esperado']} ({r['archivo']})")
        lineas.append("")
        lineas.append(f"- registros: {r['registros']}")
        lineas.append(f"- columnas: {len(r['columnas_reales'])}; identicas a lista 59: {r['columnas_identicas']}")
        if r["columnas_extra"]:
            lineas.append(f"- columnas extra no esperadas: {sorted(r['columnas_extra'].keys())}")
        lineas.append(f"- periodo unicos: {r['periodos_vistos']}")
        lineas.append("- puntajes (nulos | no_numericos | min | max | media | fuera_escala):")
        for columna, limite in COLUMNAS_PUNTAJE.items():
            a = r["puntajes"][columna]
            lineas.append(
                f"  - {columna} [0-{limite}]: {a['nulos']} | {a['no_numericos']} | "
                f"{fmt_num(a['min'])} | {fmt_num(a['max'])} | {fmt_num(a['media'])} | {a['fuera_escala']}"
            )
        lineas.append(
            f"- {COLUMNA_DEPTO}: nulos={r['depto_nulos']}, nunique={len(r['depto_unicos'])}, "
            f"formato={r['depto_formato']}, codigos={r['depto_unicos']}"
        )
        lineas.append(f"- {COLUMNA_MCPio}: nunique={r['mcpio_nunique']}")
        nulos = {c: n for c, n in r["nulos_por_columna"].items() if n > 0}
        lineas.append(f"- columnas con nulos>0: {len(nulos)} de 59")
        for c, n in sorted(nulos.items(), key=lambda x: -x[1]):
            marca = "  <-- cole_bilingue" if c == "cole_bilingue" else ""
            lineas.append(f"  - {c}: {n}{marca}")
        lineas.append(f"- duplicados exactos: {r['duplicados']}")
        if "ejemplos_duplicados" in r:
            lineas.append("- ejemplos de duplicados (filas completas con ';'):")
            lineas.append("  ```")
            lineas.append(f"  {SEPARADOR.join(COLUMNAS_ESPERADAS_59)}")
            for ej in r["ejemplos_duplicados"]:
                lineas.append(f"  {ej}")
            lineas.append("  ```")
        lineas.append("- diagnostico Fase 1 (SIN aplicar):")
        lineas.append(f"  - depto nulo: {r['diag_depto_nulo']}")
        lineas.append(f"  - punt_global == 0: {r['diag_global_cero']}")
        lineas.append(f"  - punt_global nulo: {r['diag_global_nulo']}")
        lineas.append(f"  - alguna de las 5 areas nula: {r['diag_alguna_area_nula']}")
        lineas.append("")
    lineas.append("## D — Checkpoints contra el documento (expectativas, no resultados)")
    lineas.append("")
    diferencia = total_registros - EXPECTATIVA_TOTAL_REGISTROS
    lineas.append(f"- total medido: {total_registros}; expectativa documento: {EXPECTATIVA_TOTAL_REGISTROS}")
    lineas.append(f"- diferencia exacta (medido - expectativa): {diferencia}")
    p1 = sorted(
        (r["registros"], r["periodo_esperado"]) for r in resultados
        if r["periodo_esperado"].endswith("1")
    )
    lineas.append(f"- registros de periodos P1 existentes (min-max): {p1[0][0]}-{p1[-1][0]} {p1}")
    lineas.append("- hipotesis 'falta el periodo 20181': ver bitacora SESION-002.")
    lineas.append("- referencia cruzada filtros documento (606214 | 22 | 28394) vs diagnostico medido:")
    for r in resultados:
        lineas.append(
            f"  - {r['periodo_esperado']}: depto_nulo={r['diag_depto_nulo']}, "
            f"global_cero={r['diag_global_cero']}, global_nulo={r['diag_global_nulo']}, "
            f"area_nula={r['diag_alguna_area_nula']}"
        )
    totales_diag = {
        "depto_nulo": sum(r["diag_depto_nulo"] for r in resultados),
        "global_cero": sum(r["diag_global_cero"] for r in resultados),
        "global_nulo": sum(r["diag_global_nulo"] for r in resultados),
        "area_nula": sum(r["diag_alguna_area_nula"] for r in resultados),
    }
    lineas.append(f"- totales diagnostico 11 archivos: {totales_diag}")
    lineas.append(f"- referencia documento: {EXPECTATIVA_FILTROS_REF} (sin mapeo asumido)")
    lineas.append("")
    lineas.append("## Seccion IPM")
    lineas.append("")
    lineas.append("Ver datos/output/reports/01_inventario_ipm.md.")
    lineas.append("")
    (DIR_SALIDA / "01_inventario.md").write_text("\n".join(lineas), encoding="utf-8")


def escribir_ipm_md(ipm: dict) -> None:
    """Escribe datos/output/reports/01_inventario_ipm.md (perfilado, sin transformar el xlsx)."""
    lineas = []
    lineas.append("# 01 — Inventario IPM (SESION-002, solo perfilado)")
    lineas.append("")
    lineas.append("## Hojas del anexo (dimensiones openpyxl, solo lectura)")
    lineas.append("")
    lineas.append(f"Numero de hojas: **{ipm['n_hojas']}** (SESION-001 reporto 13).")
    lineas.append("")
    lineas.append("| hoja | dimension | filas_max | columnas_max |")
    lineas.append("|---|---|---|---|")
    for h in ipm["hojas"]:
        lineas.append(f"| {h['hoja']} | {h['dimension']} | {h['filas_max']} | {h['columnas_max']} |")
    lineas.append("")
    lineas.append("## Hoja IPM_Departamentos: primeras 60 filas TAL CUAL (header=None)")
    lineas.append("")
    lineas.append(
        f"Muestra: {ipm['muestra_nfilas']} filas x {ipm['muestra_ncolumnas']} columnas "
        f"(la hoja trae {ipm['full_filas']} filas: se pidieron 60); "
        f"hoja completa: {ipm['full_filas']} filas x {ipm['full_columnas']} columnas. "
        "Campos unidos con '|' (vacio = celda vacia o NaN por celda combinada; "
        "los saltos de linea internos vienen de celdas Excel con ajuste de texto, "
        "p. ej. 'Centros poblados y rural disperso')."
    )
    lineas.append("")
    lineas.append("```")
    lineas.extend(ipm["volcado_60"])
    lineas.append("```")
    lineas.append("")
    lineas.append("## Respuestas B2 (con cifras de la lectura completa)")
    lineas.append("")
    lineas.append(f"- anios detectados en el encabezado: {ipm['anios_detectados']}")
    lineas.append(f"- fila de grupos de anio (0-based, header=None): {ipm['fila_encabezado']}")
    lineas.append(f"- filas de datos detectadas (sin titulos ni notas): {ipm['n_filas_datos']}")
    lineas.append(f"- lista de columnas detectadas ({len(ipm['columnas_detectadas'])}):")
    for j, nombre_col in enumerate(ipm["columnas_detectadas"]):
        cobertura = ipm["cobertura_por_columna"].get(j, 0)
        lineas.append(f"  - [{j}] {nombre_col!r} (valores no nulos: {cobertura})")
    lineas.append(f"- columna territorial usada (indice): {ipm['col_territorio']}")
    lineas.append(f"- columna de codigo DIVIPOLA: {ipm['col_codigo']}")
    if ipm["codigos_divipola"] is not None:
        lineas.append(f"- codigos DIVIPOLA distintos ({len(ipm['codigos_divipola'])}): {ipm['codigos_divipola']}")
    else:
        lineas.append("- codigos DIVIPOLA: NO se detecto columna de codigo (solo nombres).")
    lineas.append(f"- unidades territoriales unicas: {ipm['n_territorios_unicos']}")
    for clase, nombres in sorted(ipm["clasificacion"].items()):
        lineas.append(f"  - {clase} ({len(nombres)}): {nombres}")
    lineas.append("")
    (DIR_SALIDA / "01_inventario_ipm.md").write_text("\n".join(lineas), encoding="utf-8")


def main() -> None:
    """Ejecuta el inventario completo y guarda los tres artefactos."""
    DIR_SALIDA.mkdir(parents=True, exist_ok=True)
    archivos = sorted(DATOS_RAW.glob("Examen_Saber_11_*_limpio.csv"))
    resultados = []
    for ruta in archivos:
        print(f"[01_inventario] perfilando {ruta.name} ...", flush=True)
        resultado = perfilar_csv_icfes(ruta)
        if resultado["duplicados"] > 0:
            resultado["ejemplos_duplicados"] = extraer_ejemplos_duplicados(ruta)
        resultados.append(resultado)
        print(
            f"[01_inventario] {ruta.name}: registros={resultado['registros']} "
            f"cabecera_59={resultado['columnas_identicas']} "
            f"duplicados={resultado['duplicados']}",
            flush=True,
        )
    total_registros = sum(r["registros"] for r in resultados)
    filas = [fila_resumen_csv(r) for r in resultados]
    pd.DataFrame(filas).to_csv(DIR_SALIDA / "01_inventario_icfes.csv", index=False, encoding="utf-8")
    escribir_inventario_md(resultados, total_registros)
    print("[01_inventario] perfilando anexo IPM ...", flush=True)
    ipm = perfilar_ipm()
    escribir_ipm_md(ipm)
    print("[01_inventario] RESUMEN FINAL", flush=True)
    print(f"  archivos ICFES: {len(resultados)}", flush=True)
    print(f"  total registros: {total_registros}", flush=True)
    print(
        f"  diferencia vs expectativa documento ({EXPECTATIVA_TOTAL_REGISTROS}): "
        f"{total_registros - EXPECTATIVA_TOTAL_REGISTROS}",
        flush=True,
    )
    print(f"  hojas IPM: {ipm['n_hojas']}; anios: {ipm['anios_detectados']}", flush=True)
    print(f"  artefactos en: {DIR_SALIDA}", flush=True)


if __name__ == "__main__":
    main()
